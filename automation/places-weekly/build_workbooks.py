"""Build the weekly Places delinquency and work order workbooks.

Usage:
  python3 -I build_workbooks.py <outdir> <asof YYYY-MM-DD> --dq <delinquency.csv> [--dq-prior <prior delinquency.csv>]
                                 --wo <work_order*.csv> [<work_order*.csv> ...]
Writes <outdir>/Delinquency_By_Property_<asof>.xlsx and <outdir>/Work_Orders_By_Property_<asof>.xlsx,
and prints a JSON summary (totals, tie-out, week-over-week changes) for writing the agenda/email.
Exit 3 if a delinquency total does not tie to the AppFolio subtotal.
"""
import csv, sys, json, re, argparse
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

ORDER = ["De Cortez", "Pebble Creek", "Pikes Place"]
thin = Side(style="thin", color="000000")
hair = Side(style="hair", color="999999")
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
MONEY = '#,##0.00;(#,##0.00);"-"'


def prop_key(title):
    t = title.lower()
    for p in ORDER:
        if p.lower().split()[0] in t:
            return p
    return None


def num(s):
    s = (s or "").replace(",", "").strip()
    if not s:
        return 0.0
    neg = s.startswith("(") and s.endswith(")")
    v = float(s.strip("()"))
    return -v if neg else v


def read_dq(path):
    rows = list(csv.reader(open(path, encoding="utf-8-sig")))
    hdr = rows[0]
    ix = {h: i for i, h in enumerate(hdr)}
    out = {p: {"units": [], "subtotal": None, "title": ""} for p in ORDER}
    cur = None
    for r in rows[1:]:
        if not r or not any(c.strip() for c in r):
            continue
        if r[0].startswith("->"):
            cur = prop_key(r[0]); out[cur]["title"] = r[0][2:].strip(); continue
        if r[0] == "Total":
            continue
        if not r[0].strip():
            if cur:
                out[cur]["subtotal"] = num(r[ix["Amount Receivable"]])
            continue
        g = lambda h: r[ix[h]] if h in ix and ix[h] < len(r) else ""
        out[cur]["units"].append({
            "unit": g("Unit"), "name": g("Name"), "status": g("Tenant Status"), "tags": g("Tags"),
            "move_in": g("Move In"), "bal": num(g("Amount Receivable")), "b0": num(g("0-30")),
            "b30": num(g("30-60")), "b60": num(g("60-90")), "b90": num(g("90+")),
            "notes": " | ".join(x.strip() for x in g("Delinquency Notes").replace("\r", "").split(";\n") if x.strip()),
        })
    return out


def read_wo(paths):
    out = {p: {"wos": [], "title": ""} for p in ORDER}
    seen = set()
    for path in paths:
        rows = list(csv.reader(open(path, encoding="utf-8-sig")))
        ix = {h: i for i, h in enumerate(rows[0])}
        cur = None
        for r in rows[1:]:
            if not r or not any(c.strip() for c in r):
                continue
            if r[0].startswith("->"):
                cur = prop_key(r[0]); out[cur]["title"] = r[0][2:].strip(); continue
            g = lambda h: (r[ix[h]] if h in ix and ix[h] < len(r) else "").strip()
            n = g("Work Order Number")
            if not n or n in seen:
                continue
            p = prop_key(g("Property")) or cur
            seen.add(n)
            out[p]["wos"].append({k: g(h) for k, h in [
                ("num", "Work Order Number"), ("priority", "Priority"), ("type", "Work Order Type"), ("status", "Status"),
                ("vendor", "Vendor"), ("unit", "Unit"), ("resident", "Primary Resident"), ("created", "Created At"),
                ("sched", "Scheduled Start"), ("issue", "Work Order Issue"), ("desc", "Job Description"),
                ("instr", "Instructions")]})
    return out


def head(ws, title, sub, asof_label, asof):
    ws["A1"] = title; ws["A1"].font = F(size=12, bold=True)
    ws["A2"] = sub; ws["A2"].font = F(italic=True, color="555555")
    ws["A3"] = asof_label; ws["A3"].font = F(bold=True)
    ws["B3"] = asof; ws["B3"].number_format = "m/d/yyyy"; ws["B3"].font = F(bold=True)
    ws["B3"].alignment = Alignment(horizontal="left")


def header_row(ws, r, cols):
    for c, (name, w) in enumerate(cols, 1):
        cell = ws.cell(r, c, name)
        cell.font = F(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        cell.border = Border(top=thin, bottom=thin)
        ws.column_dimensions[L(c)].width = w
    ws.row_dimensions[r].height = 28


def dq_book(dq, prior, asof, path):
    wb = Workbook(); wb.remove(wb.active)
    summary = {}
    cols = [("Unit", 10), ("Resident", 22), ("Status", 9), ("Tags", 10), ("Move In", 11), ("Balance", 12),
            ("0-30", 11), ("30-60", 11), ("60-90", 11), ("90+", 11), ("30+", 11), ("Last Week", 12),
            ("Change", 11), ("Flag", 16), ("Delinquency Notes", 70)]
    for p in ORDER:
        d = dq[p]
        units = sorted(d["units"], key=lambda u: -u["bal"])
        pmap = {}
        if prior:
            pmap = {(u["unit"], u["name"]): u for u in prior[p]["units"]}
        ws = wb.create_sheet(p)
        head(ws, d["title"] or p, "AppFolio delinquency export. Sorted by balance, largest first." +
             (" Last Week / Change / Flag compare to the prior week's export." if prior else ""), "As of", asof)
        hr = 5
        header_row(ws, hr, cols)
        r = hr + 1
        flags = {"new": [], "up": [], "down": [], "status": [], "cleared": []}
        for u in units:
            pu = pmap.get((u["unit"], u["name"]))
            flag = ""
            if prior:
                if pu is None:
                    flag = "New"; flags["new"].append(u["unit"])
                elif pu["status"] != u["status"]:
                    flag = f"{pu['status']} -> {u['status']}"; flags["status"].append(f"{u['unit']} {flag}")
                elif u["bal"] - pu["bal"] > 0.5:
                    flag = "Increased"; flags["up"].append(u["unit"])
                elif pu["bal"] - u["bal"] > 0.5:
                    flag = "Paid down"; flags["down"].append(u["unit"])
            vals = [u["unit"], u["name"], u["status"], u["tags"],
                    datetime.strptime(u["move_in"], "%m/%d/%Y") if u["move_in"] else None,
                    u["bal"], u["b0"], u["b30"], u["b60"], u["b90"], f"=SUM(H{r}:J{r})",
                    (pu["bal"] if pu else (0 if prior else None)), (f"=F{r}-L{r}" if prior else None), flag, u["notes"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.font = F(bold=(c == 14 and bool(flag)) or (c == 3 and u["status"] == "Evict"))
                cell.border = Border(bottom=hair)
                cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 15)))
                if 6 <= c <= 13:
                    cell.number_format = MONEY
                if c == 5:
                    cell.number_format = "m/d/yyyy"
            r += 1
        last = r - 1
        ws.cell(r, 1, "Total").font = F(bold=True)
        ws.cell(r, 2, f"=COUNTA(A{hr+1}:A{last}) & \" accounts\"").font = F(bold=True)
        for c in range(6, 14):
            if c == 13 and not prior:
                continue
            if c == 12 and not prior:
                continue
            cell = ws.cell(r, c, f"=SUM({L(c)}{hr+1}:{L(c)}{last})")
            cell.font = F(bold=True); cell.number_format = MONEY
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).border = Border(top=thin, bottom=thin)
        tot = r
        r += 1
        ws.cell(r, 1, "30+ share of balance").font = F(italic=True)
        c = ws.cell(r, 6, f"=IF(F{tot}=0,0,K{tot}/F{tot})"); c.number_format = "0.0%"; c.font = F(italic=True)
        if prior:
            r += 1
            ws.cell(r, 1, "Last week's report total (incl. accounts now off the report)").font = F(italic=True)
            c = ws.cell(r, 6, round(sum(u["bal"] for u in prior[p]["units"]), 2)); c.number_format = MONEY; c.font = F(italic=True)
            ws.cell(r, 7, "Change").font = F(italic=True)
            c = ws.cell(r, 8, f"=F{tot}-F{r}"); c.number_format = MONEY; c.font = F(italic=True)
        r += 1
        ws.cell(r, 1, "AppFolio subtotal").font = F(italic=True, color="555555")
        c = ws.cell(r, 6, d["subtotal"]); c.number_format = MONEY; c.font = F(italic=True, color="555555")
        c = ws.cell(r, 7, f'=IF(ABS(F{tot}-F{r})<0.01,"ties","DOES NOT TIE")'); c.font = F(italic=True, color="555555")
        cleared = []
        if prior:
            cur_keys = {(u["unit"], u["name"]) for u in units}
            cleared = sorted([u for k, u in pmap.items() if k not in cur_keys], key=lambda u: -u["bal"])
            flags["cleared"] = [u["unit"] for u in cleared]
            if cleared:
                r += 2
                ws.cell(r, 1, "Off the report since last week (paid off or below the report filter)").font = F(bold=True)
                r += 1
                for u in cleared:
                    ws.cell(r, 1, u["unit"]).font = F(); ws.cell(r, 2, u["name"]).font = F()
                    ws.cell(r, 3, u["status"]).font = F()
                    c = ws.cell(r, 12, u["bal"]); c.number_format = MONEY; c.font = F()
                    ws.cell(r, 14, "Cleared").font = F(bold=True)
                    r += 1
        ws.freeze_panes = ws.cell(hr + 1, 3)
        ws.auto_filter.ref = f"A{hr}:{L(len(cols))}{last}"
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{hr}:{hr}"
        bal = round(sum(u["bal"] for u in units), 2)
        summary[p] = {
            "accounts": len(units), "balance": bal, "subtotal": d["subtotal"],
            "ties": d["subtotal"] is None or abs(bal - d["subtotal"]) < 0.01,
            "over30": round(sum(u["b30"] + u["b60"] + u["b90"] for u in units), 2),
            "evict": [u["unit"] for u in units if u["status"] == "Evict"],
            "notice": [u["unit"] for u in units if u["status"] == "Notice"],
            "prior_balance": round(sum(u["bal"] for u in prior[p]["units"]), 2) if prior else None,
            "flags": flags,
            "top": [{"unit": u["unit"], "status": u["status"], "bal": u["bal"], "notes": u["notes"][:200]} for u in units[:6]],
        }
    wb.save(path)
    return summary


def wo_book(wo, asof, path):
    wb = Workbook(); wb.remove(wb.active)
    summary = {}
    cols = [("Work Order #", 12), ("Priority", 9), ("Type", 10), ("Status", 10), ("Vendor", 22), ("Unit", 10),
            ("Primary Resident", 22), ("Created At", 11), ("Days Open", 8), ("Scheduled", 11), ("Issue", 20),
            ("Description", 80), ("Instructions", 28)]
    for p in ORDER:
        d = wo[p]
        wos = sorted(d["wos"], key=lambda w: (datetime.strptime(w["created"], "%m/%d/%Y"), w["num"]))
        ws = wb.create_sheet(p)
        head(ws, d["title"] or p, "AppFolio open work orders export. Sorted oldest to newest; Days Open counts to the date in B3.",
             "As of", asof)
        hr = 5
        header_row(ws, hr, cols)
        r = hr + 1
        for w in wos:
            vals = [w["num"], w["priority"], w["type"], w["status"], w["vendor"], w["unit"], w["resident"],
                    datetime.strptime(w["created"], "%m/%d/%Y"), f"=$B$3-H{r}",
                    datetime.strptime(w["sched"], "%m/%d/%Y") if w["sched"] else None, w["issue"], w["desc"], w["instr"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.font = F()
                cell.border = Border(bottom=hair)
                cell.alignment = Alignment(vertical="top", wrap_text=(c in (5, 7, 11, 12, 13)))
                if c in (8, 10):
                    cell.number_format = "m/d/yyyy"
                if c == 9:
                    cell.number_format = "0"
            r += 1
        last = r - 1
        ws.cell(r, 1, "Total").font = F(bold=True)
        ws.cell(r, 2, f"=COUNTA(A{hr+1}:A{last})" if wos else 0).font = F(bold=True)
        ws.cell(r, 8, "Avg days").font = F(bold=True)
        c = ws.cell(r, 9, f"=IFERROR(AVERAGE(I{hr+1}:I{last}),0)" if wos else 0); c.font = F(bold=True); c.number_format = "0"
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).border = Border(top=thin, bottom=thin)
        ws.freeze_panes = ws.cell(hr + 1, 2)
        if wos:
            ws.auto_filter.ref = f"A{hr}:{L(len(cols))}{last}"
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{hr}:{hr}"
        a = datetime.strptime(asof, "%Y-%m-%d")
        ages = [(a - datetime.strptime(w["created"], "%m/%d/%Y")).days for w in wos]
        summary[p] = {"count": len(wos), "avg_days": round(sum(ages) / len(ages), 1) if ages else 0,
                      "over30": sum(1 for x in ages if x > 30),
                      "oldest": [{"num": w["num"], "unit": w["unit"], "days": ag, "issue": w["issue"], "desc": w["desc"][:220]}
                                 for w, ag in list(zip(wos, ages))[:5]],
                      "all": [{"num": w["num"], "unit": w["unit"], "days": ag, "issue": w["issue"], "desc": w["desc"][:300]}
                              for w, ag in zip(wos, ages)]}
    wb.save(path)
    return summary


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("outdir"); ap.add_argument("asof")
    ap.add_argument("--dq", required=True); ap.add_argument("--dq-prior")
    ap.add_argument("--wo", nargs="+", required=True)
    a = ap.parse_args()
    asof = datetime.strptime(a.asof, "%Y-%m-%d")
    dq = read_dq(a.dq)
    prior = read_dq(a.dq_prior) if a.dq_prior else None
    dqp = f"{a.outdir}/Delinquency_By_Property_{a.asof}.xlsx"
    wop = f"{a.outdir}/Work_Orders_By_Property_{a.asof}.xlsx"
    s = {"delinquency": dq_book(dq, prior, asof, dqp), "work_orders": wo_book(read_wo(a.wo), a.asof, wop),
         "files": [dqp, wop]}
    print(json.dumps(s, indent=1, default=str))
    if not all(v["ties"] for v in s["delinquency"].values()):
        sys.exit(3)
