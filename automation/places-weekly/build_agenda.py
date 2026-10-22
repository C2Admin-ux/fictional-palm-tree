"""Render the weekly PM call agenda from a JSON spec to .docx (ink-light) and to markdown (for C2Admin call_agendas.body).

Usage: python3 -I build_agenda.py <agenda.json> <out.docx> <out.md>

agenda.json:
{
  "title": "Pikes Place, Pebble Creek & De Cortez",
  "subtitle": "Weekly Property Manager Check-In — October 28, 2026",
  "sections": [
    {"heading": "Portfolio / General",
     "subsections": [
        {"heading": "Follow-Ups from 10/21",            # optional; omit "heading" for no subheading
         "items": [
            {"lead": "Paint bids", "text": "optional one-line point",
             "subs": ["sub-point", "sub-point"]}
         ]}
     ]}
  ]
}
Each item renders as a bullet with a bold "Lead —" followed by text; subs render as indented sub-bullets.
An item with neither text nor subs renders as just the bold lead. Sections are auto-numbered.
"""
import json, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

spec = json.load(open(sys.argv[1], encoding="utf-8"))
out_docx, out_md = sys.argv[2], sys.argv[3]

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.8)
sec.top_margin = sec.bottom_margin = Inches(0.7)

BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x55, 0x55, 0x55)

st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10); st.font.color.rgb = BLACK
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(2)
for name, size in (("Heading 1", 13), ("Heading 2", 11)):
    h = doc.styles[name]
    h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = BLACK
    rf = h.element.rPr.find(qn("w:rFonts"))
    if rf is not None:
        for a in list(rf.attrib):
            del rf.attrib[a]
        rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial")
    h.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True


def rule_below(p):
    """thin bottom border under a paragraph (ink-light section divider)"""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "999999")
    bdr.append(b); pPr.append(bdr)


p = doc.add_paragraph(); r = p.add_run(spec["title"]); r.bold = True; r.font.size = Pt(15)
p = doc.add_paragraph(); r = p.add_run(spec["subtitle"]); r.italic = True; r.font.color.rgb = GREY
rule_below(p)

md = [f"# {spec['title']}", f"_{spec['subtitle']}_", ""]


def bullet(text_parts, level):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    para.paragraph_format.space_after = Pt(1)
    for t, bold in text_parts:
        run = para.add_run(t); run.bold = bold
    return para


for i, s in enumerate(spec["sections"], 1):
    doc.add_heading(f"{i}. {s['heading']}", level=1)
    md.append(f"## {i}. {s['heading']}")
    for sub in s.get("subsections", []):
        if sub.get("heading"):
            doc.add_heading(sub["heading"], level=2)
            md.append(f"### {sub['heading']}")
        items = sub.get("items", [])
        if sub.get("blank"):
            continue
        if not items:
            bullet([("Nothing open.", False)], 0)
            md.append("- Nothing open.")
        for it in items:
            lead, text, subs = it.get("lead", ""), it.get("text", ""), it.get("subs", [])
            parts = []
            if lead:
                parts.append((lead + (" — " if (text or subs) else ""), True))
            if text:
                parts.append((text, False))
            bullet(parts, 0)
            md.append("- " + (f"**{lead}**" + (" — " if (text or subs) else "") if lead else "") + text)
            for x in subs:
                bullet([(x, False)], 1)
                md.append(f"  - {x}")
    md.append("")

# blank lines for live notes under the final section
for _ in range(3):
    bullet([("", False)], 0)

# footer: source line
fp = sec.footer.paragraphs[0]
fp.text = spec.get("footer", "")
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in fp.runs:
    r.font.size = Pt(7.5); r.font.color.rgb = GREY; r.font.name = "Arial"

doc.save(out_docx)
open(out_md, "w", encoding="utf-8").write("\n".join(md).rstrip() + "\n")
print(out_docx, out_md)
